"""Builds sports-hub-deck.pptx: Arabic RTL, 16:9, native editable text, real screenshots of the trial build (demo data).
   python3 build_deck.py <crops_dir> <out.pptx>   (crops: att-record, att-team, att-month, club, lineup, player, report, roster, development)"""
import sys, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image

CROPS = pathlib.Path(sys.argv[1]); OUT = sys.argv[2]
LOGO = pathlib.Path(__file__).resolve().parents[2] / "sports-hub-app/brand/store-1024.png"
NIGHT, DEEP, INK, MUTE, PULSE, CRIM, WARN, SAFE = (RGBColor.from_string(c) for c in ("05080F", "0A1220", "EAF6FF", "93A7BB", "00E5FF", "CE1126", "FFB020", "3DDC97"))
FONT = "Tajawal"
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = 13.333, 7.5

def bg(s):
    f = s.background.fill; f.solid(); f.fore_color.rgb = NIGHT

def rect(s, x, y, w, h, fill=DEEP, line=None, radius=True):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius: sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line: sh.line.color.rgb = line; sh.line.width = Pt(1)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def para_rtl(p, align=PP_ALIGN.RIGHT):
    pPr = p._p.get_or_add_pPr(); pPr.set("rtl", "1"); p.alignment = align

def text(s, x, y, w, h, lines, size=20, color=INK, bold=False, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.TOP, space=6, bullet=False):
    """lines: str or list of str | (str, {size,color,bold}) ; **x** inside a string would be literal, so use tuples for emphasis."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor; tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, ln in enumerate(lines if isinstance(lines, list) else [lines]):
        t, o = (ln if isinstance(ln, tuple) else (ln, {}))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para_rtl(p, align); p.space_after = Pt(space)
        r = p.add_run(); r.text = (("• " if bullet else "") + t)
        f = r.font; f.name = FONT; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.color.rgb = o.get("color", color)
        rPr = r._r.get_or_add_rPr(); rPr.set("lang", "ar-BH")
        for tag in ("a:cs",):  # complex-script font so Arabic uses Tajawal too
            el = rPr.makeelement(qn(tag), {"typeface": FONT}); rPr.append(el)
    return tb

def pic(s, name, x, y, w=None, h=None, frame=True):
    path = CROPS / f"{name}.png"; im = Image.open(path); ar = im.width / im.height
    if w and not h: h = w / ar
    if h and not w: w = h * ar
    if w and h and abs(w / h - ar) > 0.01:  # fit inside the box, centred
        if w / h > ar: nw = h * ar; x += (w - nw) / 2; w = nw
        else: nh = w / ar; y += (h - nh) / 2; h = nh
    if frame: rect(s, x - 0.06, y - 0.06, w + 0.12, h + 0.12, fill=DEEP, line=RGBColor.from_string("1E3A4A"))
    p = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return p

def footer(s, n, shots=False):
    text(s, 0.4, 7.05, 6.5, 0.35, "Sports Hub · قيد التطوير · نسخة تجريبية", size=11, color=MUTE, align=PP_ALIGN.LEFT)
    if shots: text(s, 7.0, 7.05, 5.9, 0.35, "* لقطات ببيانات تجريبية: لاعبون وأندية خيالية", size=11, color=WARN)
    text(s, 12.4, 7.05, 0.6, 0.35, str(n), size=11, color=MUTE, align=PP_ALIGN.LEFT)

def title(s, t, sub=None):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.78), Inches(0.52), Inches(0.09), Inches(0.62)); bar.fill.solid(); bar.fill.fore_color.rgb = PULSE; bar.line.fill.background()
    text(s, 1.0, 0.42, 11.65, 0.9, t, size=34, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if sub: text(s, 1.0, 1.3, 11.65, 0.5, sub, size=17, color=MUTE)

def bullets(s, x, y, w, items, size=19, gap=10):
    text(s, x, y, w, 5.4, items, size=size, space=gap)

def card(s, x, y, w, h, head, body, accent=PULSE):
    rect(s, x, y, w, h, fill=DEEP, line=RGBColor.from_string("1E3A4A"))
    text(s, x + 0.2, y + 0.2, w - 0.4, 0.6, head, size=22, bold=True, color=accent)
    text(s, x + 0.2, y + 0.85, w - 0.4, h - 1.0, body, size=17, color=INK)

n = 0
def new():
    global n; n += 1; s = prs.slides.add_slide(BLANK); bg(s); return s

# 1 cover
s = new()
s.shapes.add_picture(str(LOGO), Inches(9.9), Inches(1.35), Inches(2.9), Inches(2.9))
text(s, 0.8, 1.55, 8.8, 1.1, "Sports Hub", size=60, bold=True, color=INK, align=PP_ALIGN.RIGHT)
text(s, 0.8, 2.75, 8.8, 1.3, "غرفة عمل لإدارات الأندية: الحضور والغياب، التشكيلة، تقرير المباراة، وملف اللاعب", size=26, color=PULSE)
rect(s, 6.45, 4.55, 3.15, 0.55, fill=RGBColor.from_string("2A2108"), line=WARN)
text(s, 6.45, 4.58, 3.15, 0.5, "قيد التطوير · نسخة تجريبية", size=17, bold=True, color=WARN, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.8, 5.5, 8.8, 0.9, "عرض للمناقشة وأخذ الرأي، مو عرض بيع.", size=20, color=MUTE)
footer(s, n)

# 2 questions
s = new(); title(s, "وين يضيع وقت إداري النادي؟", "ثلاثة أسئلة نبي نسمع جوابك عليها:")
for i, (h, b) in enumerate([("الحضور", "وين تسجّل حضور اللاعبين؟ وكيف تعرف مين يتكرر غيابه؟"), ("العقود والجاهزية", "متى تنتهي عقود اللاعبين؟ ومين مصاب أو موقوف قبل المباراة؟"), ("التقارير", "كيف توصّل ملخص التمرين والتقرير الشهري للمجموعة والمجلس؟")]):
    card(s, 8.95 - i * 4.1, 2.2, 3.85, 2.7, h, b)
text(s, 1.0, 5.4, 11.6, 0.8, "التطبيق مبني حول هالأسئلة. وإذا جوابك مختلف، نبي نسمعه.", size=20, color=MUTE)
footer(s, n)

# 3 idea
s = new(); title(s, "الفكرة: غرفة عمل وحدة للنادي")
s.shapes.add_picture(str(LOGO), Inches(0.9), Inches(1.7), Inches(3.6), Inches(3.6))
text(s, 0.9, 5.45, 3.6, 0.9, "حلقة (دائرة الملعب)، أربع عقد (أندية)، وكرة بالمنتصف.", size=15, color=MUTE, align=PP_ALIGN.CENTER)
for i, (h, b) in enumerate([("غرفة النادي", "«اليوم»: اللي يحتاج انتباهك الحين."), ("الحضور", "تسجيل بنقرات، نسب، جدول شهري."), ("التشكيلة والمباراة", "اقتراح تشكيلة وتقرير مباراة."), ("اللاعبون", "ملف لاعب، تطوير، حقوق وبيانات.")]):
    card(s, 5.0 + (1 - i % 2) * 3.95, 1.75 + (i // 2) * 2.45, 3.75, 2.25, h, b)
footer(s, n)

# 4 club room
s = new(); title(s, "غرفة النادي تفتح على «اليوم»", "أول شي تشوفه: اللي يحتاج قرارك.")
pic(s, "club", 0.9, 1.95, w=5.0, h=4.85)
bullets(s, 6.1, 2.1, 6.6, ["عقود تنتهي خلال 6 أشهر وعقود منتهية.", "غياب متكرر (3 غيابات بلا عذر متتالية).", "طلبات لاعبين تعدّت مهلتها الداخلية ورواتب متأخرة حسب سجلك.", "نسبة الأجور مقابل حدّها، وتقرير الموسم للاتحاد.", "كل بند يفتح المكان اللي يحلّه."])
footer(s, n, True)

# 5 attendance record
s = new(); title(s, "سجّل الحضور بنقرات، وخذ الملخص جاهز للمجموعة", "الصفحة اللي عجبت إداري جرّب النسخة.")
pic(s, "att-record", 0.9, 1.95, w=5.0, h=4.85)
bullets(s, 6.1, 2.05, 6.6, ["الكل حاضر افتراضياً، وتغيّر الاستثناءات بس.", "نقرة على اسم اللاعب تبدّل حالته: حاضر ← متأخر ← بعذر ← غائب.", "«انسخ آخر تمرين» لما تتشابه الجلسات.", "بعد الحفظ يطلع نص للمجموعة: عدد الحضور، المتأخرون والغائبون بدون عذر بأرقام القمصان.", ("ما يطلع سبب الغياب بعذر (قد يكون طبياً): العدد بس. تعدّل النص قبل ما ترسل.", {"color": PULSE})], size=18)
footer(s, n, True)

# 6 team rate
s = new(); title(s, "نسبة الفريق مقابل الـ30 يوم اللي قبلها", "وكل لاعب بنسبته، ومين يتكرر غيابه.")
pic(s, "att-team", 0.9, 1.95, w=5.0, h=4.85)
bullets(s, 6.3, 2.15, 6.4, ["الغياب بعذر ما يحسب ضد اللاعب.", "النسبة تحتاج 3 جلسات محسوبة، وإلا نكتب «ما فيه بيانات كافية».", "«غياب متكرر» = 3 غيابات بلا عذر متتالية.", "للمتابعة والكلام مع اللاعب، مو للعقاب.", "المقارنة تظهر بس لو فيه 3 جلسات بكل فترة."])
footer(s, n, True)

# 7 monthly table
s = new(); title(s, "جدول شهري يطبع بصفحة وحدة", "لاعبون × جلسات الشهر، مع نسبة كل لاعب.")
pic(s, "att-month", 0.9, 1.95, w=5.4, h=4.85)
bullets(s, 6.9, 2.15, 5.8, ["✓ حاضر · ت متأخر · ع غياب بعذر · ✗ غياب.", "زر «اطبع أو احفظ PDF» للمجلس أو المدرب.", "يطبع أبيض بدون القوائم، وتختار أي شهر."])
footer(s, n, True)

# 8 lineup + report
s = new(); title(s, "التشكيلة وتقرير المباراة", "اقتراح مسودة، والقرار للمدرب.")
pic(s, "lineup", 7.0, 1.95, w=5.6, h=4.7); pic(s, "report", 0.8, 1.95, w=5.6, h=4.7)
text(s, 7.0, 6.72, 5.7, 0.3, "تشكيلة المباراة: يقترح، المدرب يقرر", size=13, color=MUTE, align=PP_ALIGN.CENTER)
text(s, 0.8, 6.72, 5.4, 0.3, "التقرير الشهري: أرقام الفريق فقط", size=13, color=MUTE, align=PP_ALIGN.CENTER)
footer(s, n, True)

# 9 player
s = new(); title(s, "ملف اللاعب والتطوير", "حقائق وأعداد، ما فيه «درجة» للاعب.")
pic(s, "player", 7.0, 1.95, w=5.6, h=4.8); pic(s, "development", 0.8, 1.95, w=5.6, h=4.8)
footer(s, n, True)

# 10 roster
s = new(); title(s, "ضيف اللاعبين بلصق قائمة", "من واتساب أو إكسل: يقرأ الرقم والاسم والمركز لحاله.")
pic(s, "roster", 0.9, 1.95, w=5.0, h=4.85)
bullets(s, 6.1, 2.1, 6.6, ["تلصق «7 خالد المحمد مهاجم» بأي فاصل، وتشوف معاينة قبل الإضافة.", "أو تضيف لاعب لاعب بنموذج صغير.", "تُضاف للتشكيلة ولا تستبدل أحداً. بدون رواتب ولا عقود.", ("الاتفاقية أول، ومرجع الموافقة تكتبه أنت: ما نولّده.", {"color": PULSE}), "عندك صورة أو PDF؟ حدّد النص منها بجوالك وانسخه والصقه: ما نرفع صور لأي سيرفر."], size=18)
footer(s, n, True)

# 11 privacy
s = new(); title(s, "الخصوصية بقفل، مو بوعد")
rows = [("اتفاقية معالجة البيانات أول", "القاعدة نفسها ترفض أي استيراد أو إضافة بدونها."), ("موافقة لكل لاعب", "مرجع الموافقة يكتبه النادي، وسحب الموافقة يمسح بيانات اللاعب."), ("الصحة: حالة فقط", "جاهز، مشكوك، مصاب، موقوف. ما نقبل ملاحظات صحية ولا تشخيصات."), ("سجل داخلي بلا أسماء", "يسجّل من غيّر أي جدول ومتى، بدون القيم."), ("بدون أدوات تتبع خارجية", "ولا إعلانات. الأخطاء والشاشات تنعدّ على قاعدتنا بدون محتوى.")]
for i, (h, b) in enumerate(rows):
    y = 1.65 + i * 0.98
    rect(s, 1.0, y, 11.65, 0.85, fill=DEEP, line=RGBColor.from_string("1E3A4A"))
    text(s, 7.1, y + 0.1, 5.4, 0.65, h, size=20, bold=True, color=PULSE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.15, y + 0.1, 5.8, 0.65, b, size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.0, 6.55, 11.65, 0.45, "الاتفاقية والمهل القانونية تنتظر مراجعة محامٍ مرخّص قبل أي بيانات لاعبين حقيقية.", size=15, color=WARN)
footer(s, n)

# 12 honest
s = new(); title(s, "شنو شغّال اليوم وشنو لا", "بصراحة.")
card(s, 6.85, 1.85, 5.8, 4.85, "شغّال بالنسخة التجريبية", ["كل الشاشات اللي شفتها، ببيانات خيالية تنحفظ على جهازك بس.", "الحضور، الجدول الشهري، ملخص الواتساب (ينتسخ)، نسبة الفريق.", "طباعة التقرير والجدول، وضع الملعب (خط أكبر وتباين أعلى)."], accent=SAFE)
card(s, 0.7, 1.85, 5.8, 4.85, "لسا مو جاهز", ["دخول بحسابات وحفظ بياناتك على السيرفر: يحتاج إعدادات تشغيل.", "تطبيق بالمتاجر.", "إرسال تلقائي للواتساب (اليوم تنسخ النص وترسله أنت).", "قراءة صور القوائم.", "مراجعة المحامي للاتفاقية."], accent=WARN)
footer(s, n)

# 13 questions
s = new(); title(s, "رأيك يهمنا")
for i, (h, b) in enumerate([("وش أول شي تحتاجه؟", "اللي لو توفّر بالتطبيق تستخدمه بكرة."), ("شنو ما فهمته؟", "أي شاشة أو كلمة احتجت تسأل عنها."), ("شنو يضيّع وقتك اليوم؟", "وما غطّيناه: قل لنا حتى لو ما ينفع للتطبيق.")]):
    card(s, 8.95 - i * 4.1, 1.9, 3.85, 3.3, h, b)
text(s, 1.0, 5.55, 11.65, 1.2, ["ما نبيع شي. نبي نفهم شغلك قبل ما نبني أكثر.", ("وإذا ما يناسبك، قل لنا بكل أريحية.", {"color": MUTE})], size=24, align=PP_ALIGN.CENTER)
footer(s, n)

prs.save(OUT); print("saved", OUT, n, "slides")
